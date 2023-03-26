from docx import Document  # 导入库

path = ".\字典实例.docx"  # 文件路径
document = Document(path)  # 读入文件
tables01 = document.tables  # 获取文件中的表格集
table = tables01[0]  # 获取文件中的第一个表格
for i in range(1, len(table.rows)):
    result = table.cell(i, 0).text + ' ' + table.cell(i, 1).text + table.cell(i, 2).text + table.cell(i, 3).text
    print(result)
print(len(table.rows))  # 表格行数
print(len(table.columns))  # 表格列数
print(table.cell(i, 2).text)  # 表格内容
