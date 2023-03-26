from docx import Document  # 导入库

path = ".\字典实例.docx"  # 文件路径
document = Document(path)  # 读入文件
tables01 = document.tables  # 获取文件中的表格集
table = tables01[0]  # 获取文件中的第一个表格
dic = {}
for i in range(1, len(table.rows)):

    result = table.cell(i, 0).text
    result2 = table.cell(i,1).text
    dic[result] = result2

print(dic)

# dic = {'1':'1'}
# dic['2'] = '2'

# print(dic)